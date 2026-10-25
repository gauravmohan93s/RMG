from types import SimpleNamespace

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.analysis import signals as sig
from src.analysis.seasonality import productivity_lookup
from src.reports.forms import collect_forms, make_forms


def _res():
    months = pd.date_range("2025-10-01", periods=12, freq="MS")
    hist = pd.DataFrame([{"Level": "Region", "Country": "UK", "Region": r, "Metric": m, "Date": d, "Actual": a}
                         for r, a in (("North", 120.0), ("South", 60.0)) for m in ("Received", "Submitted")
                         for d in months])
    fdates = pd.date_range("2026-10-01", periods=12, freq="MS")
    fc = pd.DataFrame([{"Level": lvl, "Country": "UK", "Region": reg, "Metric": m, "Date": d, "Horizon": h + 1,
                        "Forecast": 100.0, "Lo80": 80.0, "Hi80": 120.0}
                       for lvl, reg in (("Country", "_TOTAL"), ("Region", "North"), ("Region", "South"))
                       for m in ("Received", "Submitted") for h, d in enumerate(fdates)])
    tot = pd.DataFrame([{"Level": "Country", "Country": "UK", "Metric": "Received", "Label": "United Kingdom"}])
    return SimpleNamespace(history=hist, forecast=fc, totals=tot, last_actual=pd.Timestamp("2026-09-01"))


def test_forms_round_trip_into_settings(tmp_path):
    folder = tmp_path / "Forms" / "2026-10"
    files = make_forms(_res(), folder)
    f = files["United Kingdom"]
    staff = pd.read_excel(f, "Staffing")
    assert set(staff["Region"]) == {"North", "South"} and staff["Avg_Monthly_Last_12m"].max() == 120
    inp = pd.read_excel(f, "Market_Inputs", dtype={"Month": str})
    assert len(inp) == 6 + 3 + 3            # country total x 2 metrics x 3 months + 2 regions x 3 (Received)

    # a country lead fills it in: one direct rate, one headcount only, one market input
    staff = staff.astype({"Owner": object, "Notes": object})
    inp = inp.astype({"Submitted_By": object, "Reason": object, "Submitted_On": object})
    staff.loc[staff["Region"].eq("North") & staff["Metric"].eq("Received"), ["Apps_Per_FTE_Month", "Owner"]] = [40, "A"]
    staff.loc[staff["Region"].eq("South") & staff["Metric"].eq("Received"), ["Current_FTE", "Owner"]] = [3, "A"]
    inp.loc[0, ["Expected_Volume", "Submitted_By", "Reason"]] = [130, "A", "new partner"]
    ret = folder / "Returned" / "United_Kingdom.xlsx"
    with pd.ExcelWriter(ret) as w:
        staff.to_excel(w, sheet_name="Staffing", index=False)
        inp.to_excel(w, sheet_name="Market_Inputs", index=False)

    settings = tmp_path / "Settings.xlsx"
    with pd.ExcelWriter(settings) as w:
        pd.DataFrame({"Country": ["x"]}).to_excel(w, sheet_name="Country", index=False)
    rep = collect_forms(folder, settings)
    assert rep.loc[rep["Sheet"] == "Staffing", "Accepted"].iloc[0] == 2
    assert "Country" in load_workbook(settings).sheetnames                     # other sheets kept
    prod, _ = sig.load_sheet(settings, "Staffing_Productivity", None, ["Country", "Region", "Metric", "Apps_Per_FTE_Month"])
    look = productivity_lookup(prod)
    assert look("UK", "North", "Received")[0] == 40
    assert np.isclose(look("UK", "South", "Received")[0], 60 / 3)              # derived from headcount
    hi, _ = sig.load_sheet(settings, "Human_Inputs", None, sig.HUMAN_COLS)
    h = sig.clean_human(hi)
    assert len(h) == 1 and h["Expected_Volume"].iloc[0] == 130 and h["Level"].iloc[0] == "Country"
