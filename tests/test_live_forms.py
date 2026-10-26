from types import SimpleNamespace

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.analysis import signals as sig
from src.analysis.seasonality import productivity_lookup
from src.reports.live_forms import HEADER_ROW, refresh_live, sync_live


def _res(first="2026-10-01"):
    last = pd.Timestamp(first) - pd.DateOffset(months=1)
    months = pd.date_range(last - pd.DateOffset(months=11), periods=12, freq="MS")
    hist = pd.DataFrame([{"Level": "Region", "Country": "UK", "Region": r, "Metric": m, "Date": d, "Actual": a}
                         for r, a in (("North", 120.0), ("South", 60.0)) for m in ("Received", "Submitted")
                         for d in months])
    fdates = pd.date_range(first, periods=12, freq="MS")
    fc = pd.DataFrame([{"Level": lvl, "Country": "UK", "Region": reg, "Metric": m, "Date": d, "Horizon": h + 1,
                        "Forecast": 100.0, "Lo80": 80.0, "Hi80": 120.0}
                       for lvl, reg in (("Country", "_TOTAL"), ("Region", "North"), ("Region", "South"))
                       for m in ("Received", "Submitted") for h, d in enumerate(fdates)])
    tot = pd.DataFrame([{"Level": "Country", "Country": "UK", "Metric": "Received", "Label": "United Kingdom"}])
    return SimpleNamespace(history=hist, forecast=fc, totals=tot, last_actual=last)


def _cell(ws, region, metric, row, header):
    cols = {ws.cell(HEADER_ROW + 1, c).value: c for c in range(1, ws.max_column + 1)}
    for r in range(HEADER_ROW + 2, ws.max_row + 1):
        if (ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value) == (region, metric, row):
            return ws.cell(r, cols[header])
    raise KeyError((region, metric, row))


def test_live_sheet_grid_round_trip_and_roll_forward(tmp_path):
    live = tmp_path / "Live"
    log = refresh_live(_res(), live)
    assert list(log["Status"]) == ["created"]
    f = live / "United_Kingdom.xlsx"
    wb = load_workbook(f)
    assert wb.sheetnames[:3] == ["Market View", "Staffing", "Read me"] and wb["_key"].sheet_state == "hidden"
    ws = wb["Market View"]
    assert _cell(ws, "North", "Received", "Model", "Oct 2026").value == 100
    assert _cell(ws, "North", "Received", "Actual", "Sep 2026").value == 120
    assert ws.protection.sheet                                           # model rows locked
    assert _cell(ws, "North", "Received", "Your view", "Oct 2026").protection.locked is False

    # the team fills it in place: a view on the total and on a region, and the headcount of one region
    _cell(ws, "All regions (total)", "Received", "Your view", "Nov 2026").value = 260
    _cell(ws, "North", "Received", "Your view", "Oct 2026").value = 130
    _cell(ws, "North", "Received", "Your view", "Reason / driver").value = "new partner"
    _cell(ws, "North", "Received", "Your view", "Name").value = "A"
    st = wb["Staffing"]
    south = next(r for r in range(HEADER_ROW + 2, st.max_row + 1) if st.cell(r, 1).value == "South")
    st.cell(south, 4).value = 3
    wb.save(f)

    settings = tmp_path / "Settings.xlsx"
    with pd.ExcelWriter(settings) as w:
        pd.DataFrame({"Country": ["x"]}).to_excel(w, sheet_name="Country", index=False)
        pd.DataFrame({"Country": ["Canada", "UK"], "Region": ["", ""], "Metric": ["Received"] * 2, "Month": ["2026-10"] * 2,
                      "Expected_Volume": [5, 999], "Submitted_By": ["B", "old"], "Reason": ["", ""],
                      "Submitted_On": ["2026-10-01"] * 2}).to_excel(w, sheet_name="Human_Inputs", index=False)
    rep = sync_live(live, settings)
    assert rep["Views"].iloc[0] == 2 and rep["FTE_Rows"].iloc[0] == 1
    assert "Country" in load_workbook(settings).sheetnames                  # other sheets kept
    hi, _ = sig.load_sheet(settings, "Human_Inputs", None, sig.HUMAN_COLS)
    h = sig.clean_human(hi)
    assert set(h["Country"]) == {"UK", "Canada"} and 999 not in set(h["Expected_Volume"])   # UK replaced, Canada kept
    north = h[h["Region"] == "North"].iloc[0]
    assert north["Expected_Volume"] == 130 and north["Submitted_By"] == "A" and north["Reason"] == "new partner"
    total = h[(h["Level"] == "Country") & (h["Country"] == "UK")].iloc[0]
    assert total["Expected_Volume"] == 260 and total["Submitted_By"] == "United Kingdom team"
    prod, _ = sig.load_sheet(settings, "Staffing_Productivity", None, ["Country", "Region", "Metric", "Apps_Per_FTE_Month"])
    assert np.isclose(productivity_lookup(prod)("UK", "South", "Received")[0], 60 / 3)

    # next month: the window rolls forward, the entries stay (Oct becomes a past, locked month)
    log = refresh_live(_res("2026-11-01"), live)
    assert log["Status"].iloc[0] == "refreshed" and (live / "_history").exists()
    ws = load_workbook(f)["Market View"]
    oct_view = _cell(ws, "North", "Received", "Your view", "Oct 2026")
    assert oct_view.value == 130 and oct_view.protection.locked is not False
    assert _cell(ws, "All regions (total)", "Received", "Your view", "Nov 2026").value == 260
    assert _cell(ws, "North", "Received", "Your view", "Name").value == "A"
    assert load_workbook(f)["Staffing"].cell(south, 4).value == 3


def test_unreadable_live_sheet_is_never_overwritten(tmp_path):
    live = tmp_path / "Live"
    live.mkdir()
    bad = live / "United_Kingdom.xlsx"
    bad.write_bytes(b"not an excel file")
    log = refresh_live(_res(), live)
    assert log["Status"].iloc[0] == "skipped" and bad.read_bytes() == b"not an excel file"
